import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import List, Dict, Any
from pathlib import Path

def generate_candidates_excel(candidates: List[Dict[str, Any]], output_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Arizalar Ro'yxati"

    # Sarlavha uslubi
    title_font = Font(name="Calibri", size=15, bold=True, color="1F497D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    status_colors = {
        "submitted": ("FFF2CC", "D6B656"),   # Sariq / Kutilmoqda
        "interview": ("D5E8D4", "82B366"),   # Yashil / Suhbat
        "reserve": ("FFE6CC", "D79B00"),     # To'q sariq / Zaxira
        "rejected": ("F8CECC", "B85450")     # Qizil / Rad etilgan
    }

    # Asosiy sarlavha qatori
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = f"Ixtisoslashgan maktab kadrlar tanlovi bo'yicha nomzodlar ro'yxati (Yuklangan sana: {datetime.now().strftime('%d.%m.%Y %H:%M')})"
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Ustun sarlavhalari
    headers = [
        "№",
        "Ariza ID",
        "F.I.SH",
        "Telefon raqam",
        "Vakansiya",
        "Toifasi",
        "Ish staji",
        "Ma'lumoti / OTM",
        "Sertifikatlari",
        "Holati",
        "Topshirilgan sana",
        "Amal bajargan xodim",
        "Amal vaqti"
    ]

    ws.row_dimensions[3].height = 25
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_labels = {
        "submitted": "Yangi (Kutilmoqda)",
        "interview": "Suhbatga chaqirilgan",
        "reserve": "Zaxiraga olingan",
        "rejected": "Rad etilgan"
    }

    # Ma'lumotlarni yozish
    row_num = 4
    for idx, cand in enumerate(candidates, 1):
        ws.row_dimensions[row_num].height = 20
        status_key = cand.get("status", "submitted")
        status_text = status_labels.get(status_key, status_key)
        
        row_values = [
            idx,
            f"#{cand.get('id', '')}",
            cand.get("full_name", "-"),
            cand.get("phone", "-"),
            cand.get("vacancy_title", "-"),
            cand.get("category", "-"),
            cand.get("experience", "-"),
            cand.get("education", "-"),
            cand.get("cert_info", "-"),
            status_text,
            cand.get("submitted_at", "-"),
            cand.get("action_by_name", "-") or "-",
            cand.get("action_at", "-") or "-"
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center" if col_idx in [1, 2, 4, 6, 7, 10, 11, 13] else "left",
                vertical="center"
            )
            
            # Holat ustuni uchun maxsus rang berish
            if col_idx == 10 and status_key in status_colors:
                bg_color, _ = status_colors[status_key]
                cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        row_num += 1

    # Ustun kengliklarini avtomatik moslash
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Sarlavha qatorini hisobga olmaslik
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
